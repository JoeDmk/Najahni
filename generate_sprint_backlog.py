"""
Generate Sprint Backlog Excel — Gestion des Utilisateurs (Najahni)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from datetime import date, timedelta

wb = Workbook()
ws = wb.active
ws.title = "Sprint Backlog - Users"

# ── Colours ──────────────────────────────────────────────────
DARK_BLUE   = "1B2A4A"
MED_BLUE    = "2E86C1"
LIGHT_BLUE  = "D6EAF8"
WHITE       = "FFFFFF"
LIGHT_GRAY  = "F2F3F4"
GREEN       = "27AE60"
ORANGE      = "F39C12"
RED         = "E74C3C"
YELLOW_BG   = "FEF9E7"
GREEN_BG    = "EAFAF1"
ORANGE_BG   = "FDF2E9"
RED_BG      = "FDEDEC"
BLUE_BG     = "EBF5FB"
PURPLE_BG   = "F4ECF7"

# ── Styles ───────────────────────────────────────────────────
thin_border = Border(
    left=Side(style="thin", color="BDC3C7"),
    right=Side(style="thin", color="BDC3C7"),
    top=Side(style="thin", color="BDC3C7"),
    bottom=Side(style="thin", color="BDC3C7"),
)

title_font   = Font(name="Calibri", size=18, bold=True, color=WHITE)
title_fill   = PatternFill("solid", fgColor=DARK_BLUE)
header_font  = Font(name="Calibri", size=11, bold=True, color=WHITE)
header_fill  = PatternFill("solid", fgColor=MED_BLUE)
cell_font    = Font(name="Calibri", size=10)
bold_font    = Font(name="Calibri", size=10, bold=True)
center       = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap    = Alignment(horizontal="left", vertical="center", wrap_text=True)

status_styles = {
    "Terminé":    (Font(name="Calibri", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=GREEN)),
    "En cours":   (Font(name="Calibri", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=ORANGE)),
    "À faire":    (Font(name="Calibri", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=RED)),
}

category_fills = {
    "Authentification":    PatternFill("solid", fgColor=BLUE_BG),
    "CRUD Admin":          PatternFill("solid", fgColor=GREEN_BG),
    "Profil":              PatternFill("solid", fgColor=YELLOW_BG),
    "Fonctionnalités Pro": PatternFill("solid", fgColor=PURPLE_BG),
    "Sécurité":            PatternFill("solid", fgColor=ORANGE_BG),
    "UI / UX":             PatternFill("solid", fgColor=LIGHT_BLUE),
}

# ── Title row ────────────────────────────────────────────────
ws.merge_cells("A1:I1")
title_cell = ws["A1"]
title_cell.value = "Sprint Backlog — Gestion des Utilisateurs (Najahni)"
title_cell.font = title_font
title_cell.fill = title_fill
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

# ── Subtitle row ─────────────────────────────────────────────
ws.merge_cells("A2:I2")
sub = ws["A2"]
sub.value = "PIDEV 3A · Module : Gestion Users · Équipe Najahni"
sub.font = Font(name="Calibri", size=11, italic=True, color=DARK_BLUE)
sub.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
sub.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 25

# ── Headers ──────────────────────────────────────────────────
headers = ["#", "Catégorie", "Tâche", "Statut", "Owner (Email)",
           "Sprint", "Date Limite", "Notes", "User Story"]
col_widths = [5, 22, 45, 12, 30, 10, 14, 40, 50]

for col_idx, (h, w) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=3, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = thin_border
    ws.column_dimensions[get_column_letter(col_idx)].width = w
ws.row_dimensions[3].height = 28

# ── Sprint dates ─────────────────────────────────────────────
S1_START = date(2025, 2, 3)
S1_END   = date(2025, 2, 16)
S2_START = date(2025, 2, 17)
S2_END   = date(2025, 3, 2)
S3_START = date(2025, 3, 3)
S3_END   = date(2025, 3, 16)

OWNER = "ilyes@najahni.tn"

# ── Data ─────────────────────────────────────────────────────
tasks = [
    # Sprint 1 — Core auth + CRUD
    ("Authentification", "Conception de la BDD (table user, login_history, notifications)", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=2),
     "MySQL · najahni_db", "En tant qu'administrateur, je veux une base de données structurée pour stocker les utilisateurs."),

    ("Authentification", "Implémentation du modèle User (entité JPA)", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=3),
     "Classe models/User.java", "En tant que développeur, je veux un modèle User avec tous les attributs nécessaires."),

    ("Authentification", "Inscription avec validation (email, tél, mdp)", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=5),
     "SignUpController + ValidationService", "En tant qu'utilisateur, je veux m'inscrire avec un formulaire validé pour créer mon compte."),

    ("Authentification", "Vérification email par code (SMTP Gmail)", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=6),
     "EmailService · Jakarta Mail", "En tant qu'utilisateur, je veux recevoir un code par email pour vérifier mon compte."),

    ("Authentification", "Connexion par mot de passe + BCrypt", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=7),
     "SignInController + UserService", "En tant qu'utilisateur, je veux me connecter avec email et mot de passe sécurisé."),

    ("Authentification", "Connexion Google OAuth 2.0", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=9),
     "GoogleOAuthService + WebView", "En tant qu'utilisateur, je veux me connecter via Google pour un accès rapide."),

    ("Authentification", "CAPTCHA (Math Challenge) sur connexion", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=10),
     "Empêcher les bots", "En tant qu'administrateur, je veux un CAPTCHA pour protéger la page de connexion."),

    ("Authentification", "Mot de passe oublié (code email + reset)", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=11),
     "PasswordResetService", "En tant qu'utilisateur, je veux réinitialiser mon mot de passe via un code envoyé par email."),

    ("CRUD Admin", "Dashboard Admin — liste des utilisateurs", "Terminé", OWNER, "Sprint 1", S1_START + timedelta(days=12),
     "DashboardController + TableView", "En tant qu'admin, je veux voir la liste de tous les utilisateurs dans un tableau."),

    ("CRUD Admin", "Ajouter un utilisateur (admin)", "Terminé", OWNER, "Sprint 1", S1_END - timedelta(days=1),
     "AjouterUser.fxml", "En tant qu'admin, je veux ajouter un utilisateur manuellement."),

    ("CRUD Admin", "Modifier un utilisateur (admin)", "Terminé", OWNER, "Sprint 1", S1_END,
     "ModifierUser.fxml", "En tant qu'admin, je veux modifier les informations d'un utilisateur."),

    ("CRUD Admin", "Supprimer un utilisateur (admin)", "Terminé", OWNER, "Sprint 1", S1_END,
     "Confirmation dialog", "En tant qu'admin, je veux supprimer un utilisateur avec confirmation."),

    # Sprint 2 — Advanced features
    ("CRUD Admin", "Bannir / Débannir un utilisateur", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=1),
     "is_banned toggle", "En tant qu'admin, je veux bannir/débannir un utilisateur pour gérer les accès."),

    ("CRUD Admin", "Recherche et filtrage (nom, email, rôle)", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=2),
     "SearchField + ComboBox rôle", "En tant qu'admin, je veux rechercher et filtrer les utilisateurs par critères."),

    ("Profil", "Page Profil utilisateur (consultation)", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=3),
     "Profil.fxml + ProfilController", "En tant qu'utilisateur, je veux consulter mon profil avec mes informations."),

    ("Profil", "Modification du profil utilisateur", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=4),
     "Validation + updateUser", "En tant qu'utilisateur, je veux modifier mes informations personnelles."),

    ("Sécurité", "Face ID — Enregistrement (20 échantillons)", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=7),
     "OpenCV LBPH · 20 samples · model.yml", "En tant qu'utilisateur, je veux enregistrer mon visage pour la connexion Face ID."),

    ("Sécurité", "Face ID — Connexion par reconnaissance faciale", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=9),
     "FaceLoginController · threshold 80 · 5 matchs", "En tant qu'utilisateur, je veux me connecter par reconnaissance faciale."),

    ("Sécurité", "Face ID — Persistance du modèle global", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=10),
     "global_model.yml sauvegardé sur disque", "En tant qu'utilisateur, je veux que le Face ID fonctionne après redémarrage du PC."),

    ("Sécurité", "Détection de connexion suspecte (IA)", "Terminé", OWNER, "Sprint 2", S2_START + timedelta(days=11),
     "SuspiciousLoginService · score risque", "En tant qu'admin, je veux détecter les connexions suspectes automatiquement."),

    ("Sécurité", "Historique de connexion", "Terminé", OWNER, "Sprint 2", S2_END - timedelta(days=2),
     "LoginHistoryService · tableau 100 derniers", "En tant qu'admin, je veux voir l'historique des connexions (méthode, IP, appareil)."),

    ("Sécurité", "Verrouillage après tentatives échouées", "Terminé", OWNER, "Sprint 2", S2_END,
     "SessionService · isAccountLocked", "En tant qu'admin, je veux verrouiller un compte après trop de tentatives."),

    # Sprint 3 — Pro features + polish
    ("Fonctionnalités Pro", "Export CSV (OpenCSV)", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=1),
     "ExportService · OpenCSV", "En tant qu'admin, je veux exporter la liste des utilisateurs en CSV."),

    ("Fonctionnalités Pro", "Export PDF (iText 7)", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=2),
     "ExportService · iText7", "En tant qu'admin, je veux exporter la liste des utilisateurs en PDF."),

    ("Fonctionnalités Pro", "Email Broadcast (envoi groupé)", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=4),
     "EmailService · SMTP à tous", "En tant qu'admin, je veux envoyer un email groupé à tous les utilisateurs."),

    ("Fonctionnalités Pro", "Notifications (cloche + marquer lu)", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=5),
     "NotificationService · badge 🔔", "En tant qu'utilisateur, je veux recevoir et consulter des notifications."),

    ("UI / UX", "Thème Sombre / Clair (toggle)", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=7),
     "ThemeService · dark-theme.css", "En tant qu'utilisateur, je veux basculer entre thème sombre et clair."),

    ("Fonctionnalités Pro", "Statistiques (PieChart + BarChart)", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=9),
     "StatsView.fxml · JavaFX Charts", "En tant qu'admin, je veux visualiser les statistiques des utilisateurs."),

    ("Fonctionnalités Pro", "Auto-suggestion utilisateurs", "Terminé", OWNER, "Sprint 3", S3_START + timedelta(days=10),
     "Suggestions dans la recherche", "En tant qu'admin, je veux des suggestions automatiques lors de la recherche."),

    ("UI / UX", "Design moderne (CSS najahni.css, modern.css)", "Terminé", OWNER, "Sprint 3", S3_END - timedelta(days=2),
     "JavaFX CSS · responsive", "En tant qu'utilisateur, je veux une interface moderne et ergonomique."),

    ("UI / UX", "Tests et corrections de bugs", "Terminé", OWNER, "Sprint 3", S3_END,
     "Face ID fix · tooltip fix · stats fix", "En tant que développeur, je veux corriger tous les bugs identifiés."),

    ("UI / UX", "Documentation (diagramme de séquence UML)", "Terminé", OWNER, "Sprint 3", S3_END,
     "PlantUML · 3 couches · 16 flux", "En tant que développeur, je veux documenter l'architecture avec un diagramme UML."),
]

# ── Write data rows ──────────────────────────────────────────
for i, (cat, task, status, owner, sprint, due, notes, story) in enumerate(tasks, 1):
    row = i + 3  # header is row 3

    values = [i, cat, task, status, owner, sprint, due, notes, story]

    # Alternate row background
    cat_fill = category_fills.get(cat, PatternFill("solid", fgColor=LIGHT_GRAY))
    alt_fill = PatternFill("solid", fgColor=LIGHT_GRAY) if i % 2 == 0 else PatternFill("solid", fgColor=WHITE)

    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = cell_font
        cell.border = thin_border

        # Column-specific formatting
        if col_idx == 1:   # #
            cell.alignment = center
            cell.fill = alt_fill
        elif col_idx == 2: # Category
            cell.alignment = center
            cell.fill = cat_fill
            cell.font = bold_font
        elif col_idx == 3: # Task
            cell.alignment = left_wrap
            cell.fill = alt_fill
        elif col_idx == 4: # Status
            cell.alignment = center
            s_font, s_fill = status_styles.get(status, (cell_font, alt_fill))
            cell.font = s_font
            cell.fill = s_fill
        elif col_idx == 5: # Owner
            cell.alignment = center
            cell.fill = alt_fill
        elif col_idx == 6: # Sprint
            cell.alignment = center
            cell.fill = alt_fill
            cell.font = bold_font
        elif col_idx == 7: # Due date
            cell.alignment = center
            cell.number_format = "DD/MM/YYYY"
            cell.fill = alt_fill
        elif col_idx == 8: # Notes
            cell.alignment = left_wrap
            cell.fill = alt_fill
        elif col_idx == 9: # User Story
            cell.alignment = left_wrap
            cell.fill = alt_fill

    ws.row_dimensions[row].height = 32

# ── Summary section ──────────────────────────────────────────
summary_row = len(tasks) + 5
ws.merge_cells(f"A{summary_row}:B{summary_row}")
sc = ws.cell(row=summary_row, column=1, value="Résumé du Sprint Backlog")
sc.font = Font(name="Calibri", size=13, bold=True, color=DARK_BLUE)
sc.alignment = Alignment(horizontal="left", vertical="center")

labels = [
    ("Total tâches :", len(tasks)),
    ("Terminé :", sum(1 for t in tasks if t[2] == "Terminé")),
    ("En cours :", sum(1 for t in tasks if t[2] == "En cours")),
    ("À faire :", sum(1 for t in tasks if t[2] == "À faire")),
]

for j, (lbl, val) in enumerate(labels):
    r = summary_row + 1 + j
    lc = ws.cell(row=r, column=1, value=lbl)
    lc.font = bold_font
    lc.alignment = Alignment(horizontal="right", vertical="center")
    vc = ws.cell(row=r, column=2, value=val)
    vc.font = Font(name="Calibri", size=11, bold=True,
                   color=GREEN if "Terminé" in lbl else (ORANGE if "cours" in lbl else (RED if "faire" in lbl else DARK_BLUE)))
    vc.alignment = center

# ── Legend ────────────────────────────────────────────────────
legend_row = summary_row + 6
ws.cell(row=legend_row, column=1, value="Légende Catégories :").font = bold_font
for k, (cat_name, fill) in enumerate(category_fills.items()):
    r = legend_row + 1 + k
    cell = ws.cell(row=r, column=1, value=cat_name)
    cell.fill = fill
    cell.font = cell_font
    cell.border = thin_border
    cell.alignment = center

# ── Freeze panes ─────────────────────────────────────────────
ws.freeze_panes = "A4"

# ── Auto-filter ──────────────────────────────────────────────
ws.auto_filter.ref = f"A3:I{len(tasks) + 3}"

# ── Print setup ──────────────────────────────────────────────
ws.sheet_properties.pageSetUpPr = None
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1

# ── Save ─────────────────────────────────────────────────────
output = r"c:\Users\ilyes\Downloads\USER\PIDEV-3A-JAVA-Gestion-Users\Sprint_Backlog_Gestion_Users.xlsx"
wb.save(output)
print(f"✅ Sprint Backlog saved to: {output}")
